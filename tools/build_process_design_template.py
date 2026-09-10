#!/usr/bin/env python3
"""認証実装確認処理 — 処理設計書（Excel）雛形ジェネレータ

SSOT: doc/api-platform/basic-design/research/process-design-template.md
出力: doc/excel/apipf-process-design.xlsx

PROCESSES を更新したら上記 md の処理カタログも同時に更新すること。
再実行すると出力ファイルを作り直す（記入済みの内容は失われるので注意）。
"""
from __future__ import annotations

import pathlib

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = pathlib.Path(__file__).resolve().parents[1] / "doc" / "excel" / "apipf-process-design.xlsx"

# ---------------------------------------------------------------- styles
TITLE = Font(bold=True, size=14, color="FFFFFF")
TITLE_FILL = PatternFill("solid", fgColor="2F5597")
SEC = Font(bold=True, size=11, color="FFFFFF")
SEC_FILL = PatternFill("solid", fgColor="8EA9DB")
HDR = Font(bold=True, size=10)
HDR_FILL = PatternFill("solid", fgColor="D9E1F2")
LBL_FILL = PatternFill("solid", fgColor="F2F2F2")
TODO_FILL = PatternFill("solid", fgColor="FFF2CC")
BASE = Font(size=10)
LINK = Font(size=10, color="0563C1", underline="single")
WRAP = Alignment(vertical="top", wrap_text=True)
TOP = Alignment(vertical="top")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

# ---------------------------------------------------------------- data
# (id, 短縮シート名, 処理名, 区分, 起動元, 起動先, 方式, 実行主体, 頻度, 概要,
#  権限/エンドポイント, シーケンス, 手順ヒント)
P = [
    ("P-01", "巡回起動", "巡回起動", "A 巡回・発見",
     "EventBridge Scheduler（rate(1 hour)）", "発見 Lambda", "非同期", "EventBridge Scheduler", "1 時間毎",
     "監視対象の変更有無を確認する巡回処理を定期起動する。",
     "Scheduler 実行ロール / lambda:InvokeFunction / lambda.{region}.amazonaws.com",
     "Scheduler -> 発見 Lambda : Invoke（payload なし）",
     "リトライポリシー・Scheduler DLQ の設定値を明記する（17 章 / WBS A6-d）"),
    ("P-02", "対象アカウント列挙", "対象アカウント列挙", "A 巡回・発見",
     "発見 Lambda", "AWS Organizations", "同期", "発見 Lambda", "巡回ごとに 1 回",
     "巡回対象となる App アカウントの一覧を取得する。",
     "organizations:ListAccounts（委任ポリシー方式が推奨。M-Q-17-2）/ organizations.us-east-1.amazonaws.com",
     "発見 Lambda -> Organizations : ListAccounts\nOrganizations --> 発見 Lambda : アカウント一覧",
     "列挙方式（委任 / AssumeRole / 静的リスト）と対象範囲（全体 / OU）を確定する"),
    ("P-03", "読み取り権限取得", "読み取り権限取得", "A 巡回・発見",
     "発見 Lambda", "STS →（各 App の）DiscoveryReadRole", "同期", "発見 Lambda", "アカウントごとに 1 回",
     "App アカウントの資材を読むための一時クレデンシャルを取得する。",
     "sts:AssumeRole + ExternalId / sts.{region}.amazonaws.com（リージョナル STS）",
     "発見 Lambda -> STS : AssumeRole(DiscoveryReadRole, ExternalId)\nSTS --> 発見 Lambda : 一時クレデンシャル",
     "クレデンシャルの有効期限とキャッシュ方針を明記する"),
    ("P-04", "資材一覧・版数取得", "資材一覧・版数取得", "A 巡回・発見",
     "発見 Lambda", "App アカウントの資材バケット", "同期", "発見 Lambda", "アカウントごとに 1 回",
     "資材バケットの {appId}/ を列挙し、各資材の VersionId を取得する（本文は取得しない）。",
     "s3:ListBucket / s3:ListBucketVersions / s3.{region}.amazonaws.com",
     "発見 Lambda -> 資材バケット : List（{appId}/ プレフィックス）\n資材バケット --> 発見 Lambda : キー一覧 + VersionId",
     "ページング処理と、monitoring.yaml が無いプレフィックスの扱いを明記する"),
    ("P-05", "台帳読取", "台帳読取", "A 巡回・発見",
     "発見 Lambda", "Monitoring Registry S3（registry/）", "同期", "発見 Lambda", "アプリごとに 1 回",
     "前回巡回時の状態（lastArtifactVersions 等）を台帳から読み取る。",
     "s3:GetObject（同一アカウント）/ s3.{region}.amazonaws.com",
     "発見 Lambda -> registry/ : GetObject(registry/{appId}/{env}.json)\nregistry/ --> 発見 Lambda : 台帳 JSON（無ければ新規扱い）",
     "台帳が存在しない場合（初回）の扱いを明記する"),
    ("P-06", "差分判定", "差分判定", "A 巡回・発見",
     "発見 Lambda（内部）", "—", "内部", "発見 Lambda", "アプリごとに 1 回",
     "資材の VersionId と台帳の lastArtifactVersions を比較し、検査要否を判定する。",
     "—（AWS 呼び出しなし）",
     "発見 Lambda : VersionId ↔ lastArtifactVersions を比較\n判定結果 = 変更あり / 変更なし / 新規",
     "【重要】比較はメタデータ（VersionId）で行い、内容ハッシュ比較にしない（17 §17.2.1 差分判定の 3 原則）"),
    ("P-07", "資材取得", "資材取得", "A 巡回・発見",
     "発見 Lambda", "App アカウントの資材バケット", "同期", "発見 Lambda", "変更のあったアプリのみ",
     "変更が検知された資材（monitoring.yaml / openapi.yaml / deploy-info.json）を取得する。",
     "s3:GetObject / s3:GetObjectVersion / s3.{region}.amazonaws.com",
     "発見 Lambda -> 資材バケット : GetObject(monitoring.yaml, openapi.yaml)\n資材バケット --> 発見 Lambda : ファイル本体",
     "サイズ上限・文字コード・取得失敗時の扱いを明記する"),
    ("P-08", "資材検証", "資材検証", "A 巡回・発見",
     "発見 Lambda（内部）", "—", "内部", "発見 Lambda", "取得ごとに 1 回",
     "monitoring.yaml のスキーマ検証と appId／プレフィックス一致検証を行う。",
     "—（AWS 呼び出しなし）",
     "発見 Lambda : YAML パース -> JSON Schema 検証 -> appId 一致検証\nNG の場合は P-12 メタ不足通知へ",
     "不備の分類（取り込み拒否 / 既定値で継続）を項目ごとに定義する（17 §17.3 の表）"),
    ("P-09", "台帳登録・更新", "台帳登録・更新", "A 巡回・発見",
     "発見 Lambda", "Monitoring Registry S3（registry/）", "同期", "発見 Lambda", "変更のあったアプリのみ",
     "台帳へ設定値を同期し、巡回状態を更新する。中央管理項目は上書きしない。",
     "s3:PutObject（ETag 条件付き PUT, If-Match）/ s3.{region}.amazonaws.com",
     "発見 Lambda -> registry/ : PutObject（If-Match: 読取時 ETag）\n412 の場合は再読取してリトライ",
     "alertRouting / enabled を上書きしないマージ規則と、412 競合時のリトライ回数を明記する"),
    ("P-10", "spec 配置", "spec 配置", "A 巡回・発見",
     "発見 Lambda", "Monitoring Registry S3（openapi/）", "同期", "発見 Lambda", "変更のあったアプリのみ",
     "取得した openapi.yaml を中央の spec 置き場へコピーする。",
     "s3:PutObject / s3.{region}.amazonaws.com",
     "発見 Lambda -> openapi/ : PutObject(openapi/{accountId}/{appId}/openapi.yaml)",
     "キー導出規則と上書き方針（Versioning による履歴保全）を明記する"),
    ("P-11", "検査起動", "検査起動", "A 巡回・発見",
     "発見 Lambda", "認証実装チェック Lambda", "非同期", "発見 Lambda", "変更のあったアプリごと",
     "変更が確定したアプリを対象に検査を起動する（自動差分検査／モード1）。",
     "lambda:InvokeFunction（InvocationType=Event）/ lambda.{region}.amazonaws.com",
     "発見 Lambda -> チェック Lambda : Invoke(Event, {mode:'delta', appId, env})\n※ 起動成功後に P-09 の lastArtifactVersions を確定更新",
     "非同期リトライ（2 回）と DLQ の設定、起動失敗時に台帳を更新しないこと（at-least-once）を明記する"),
    ("P-12", "メタ不足通知", "メタ不足通知", "A 巡回・発見",
     "発見 Lambda", "SNS（P2 Platform）", "同期", "発見 Lambda", "不備検出時",
     "資材の不備（appId 不一致・必須項目欠落・authPattern 不正など）を通知する。",
     "sns:Publish / sns.{region}.amazonaws.com",
     "発見 Lambda -> SNS(P2) : Publish（不備種別・appId・該当項目）",
     "不備種別ごとの文面と、毎巡回で重複通知しない抑制方式を定義する"),
    ("P-13", "消滅staleness検知", "消滅・staleness 検知", "A 巡回・発見",
     "発見 Lambda", "registry/ + SNS（P2）", "同期", "発見 Lambda", "巡回ごとに 1 回",
     "資材の消滅（アプリ廃止の可能性）と、長期間更新されていない資材を検知する。",
     "s3:PutObject（enabled=false）/ sns:Publish",
     "発見 Lambda : 台帳あり かつ 資材なし -> enabled=false + 棚卸しアラート\n発見 Lambda : 最終更新が閾値超 -> staleness アラート",
     "staleness の閾値（仮 90 日、M-Q-17-3）と、一時的な取得失敗との区別を定義する"),
    ("P-14", "巡回メトリクス", "巡回結果メトリクス出力", "A 巡回・発見",
     "発見 Lambda", "CloudWatch", "同期", "発見 Lambda", "巡回ごとに 1 回",
     "巡回の成功・失敗アカウント数などのメタ監視用メトリクスを出力する。",
     "cloudwatch:PutMetricData / monitoring.{region}.amazonaws.com",
     "発見 Lambda -> CloudWatch : PutMetricData(DiscoveryLastSuccess, DiscoveryAccountErrors)",
     "メトリクス名・ディメンション・欠損時の扱い（MM-1 は 2h 欠損で発報）を明記する"),

    ("P-20", "全量検査 定期起動", "全量検査の定期起動", "B 検査",
     "EventBridge Scheduler（日次）", "認証実装チェック Lambda", "非同期", "EventBridge Scheduler", "日次",
     "資材の変化に関係なく全アプリを検査するため、全量検査を定期起動する。",
     "Scheduler 実行ロール / lambda:InvokeFunction",
     "Scheduler -> チェック Lambda : Invoke(Event, {mode:'full'})",
     "実行時刻帯（M-Q-18-3）と、巡回（P-01）と時間帯が重ならない配慮を明記する"),
    ("P-21", "全量検査 手動起動", "全量検査の手動起動", "B 検査",
     "運用者（CLI / コンソール）", "認証実装チェック Lambda", "非同期", "運用者", "随時（監査前・障害後など）",
     "運用者の判断で全量検査を実行する。実装は定期起動と同一。",
     "lambda:InvokeFunction（実行者の IAM）",
     "運用者 -> チェック Lambda : Invoke(Event, {mode:'full'})\n※ 特定アプリのみ: {mode:'full', appId, env}",
     "実行権限を持つ者の範囲（共通基盤チームのみか、アプリチームも可か。M-Q-18-3）を確定する"),
    ("P-22", "台帳取得", "台帳取得", "B 検査",
     "認証実装チェック Lambda", "Monitoring Registry S3（registry/）", "同期", "認証実装チェック Lambda", "検査ごとに 1 回",
     "検査対象の設定値を台帳から取得する。全量時は List、差分時は該当 1 件のみ Get。",
     "s3:ListBucket / s3:GetObject",
     "チェック Lambda -> registry/ : List（mode=full）または GetObject（mode=delta）\nregistry/ --> チェック Lambda : 台帳 JSON",
     "enabled=false のレコードを除外すること、List のページングを明記する"),
    ("P-23", "fan-out", "アプリ単位 fan-out", "B 検査",
     "認証実装チェック Lambda", "認証実装チェック Lambda（自身）", "非同期", "認証実装チェック Lambda", "全量検査時のみ",
     "全量検査を「1 実行 = 1 アプリ」に分割し、Lambda の実行時間上限に構造的に当たらないようにする。",
     "lambda:InvokeFunction（自身）/ lambda.{region}.amazonaws.com",
     "チェック Lambda(mode=full) -> チェック Lambda(mode=delta 相当) : アプリ数分 Invoke(Event)",
     "同時実行数の上限・スロットリング時の扱いを明記する"),
    ("P-24", "spec 取得", "spec 取得", "B 検査",
     "認証実装チェック Lambda", "Monitoring Registry S3（openapi/）", "同期", "認証実装チェック Lambda", "アプリごとに 1 回",
     "検査対象 endpoint の一覧と公開印（MON-1）を取得する。",
     "s3:GetObject",
     "チェック Lambda -> openapi/ : GetObject(openApiS3Key)\nopenapi/ --> チェック Lambda : spec",
     "spec が存在しない場合（モノリスの endpoint 列挙方式）の分岐を明記する"),
    ("P-25", "検査用トークン取得", "検査用トークン取得", "B 検査",
     "認証実装チェック Lambda", "Secrets Manager → 認証基盤 /token", "同期", "認証実装チェック Lambda", "検査ごとに 1 回",
     "正規検査（Positive probe）で使う短命トークンを取得する。",
     "secretsmanager:GetSecretValue / 認証基盤の公開 /token（インターネット経由・OAuth client_credentials）",
     "チェック Lambda -> Secrets Manager : GetSecretValue\nチェック Lambda -> Keycloak /token : client_credentials\nKeycloak --> チェック Lambda : access_token（短命）",
     "トークンのキャッシュ有無・失効時の再取得・ログへのマスク（トークンを出力しない）を明記する"),
    ("P-26", "未認証検査", "未認証検査（Negative probe）", "B 検査",
     "認証実装チェック Lambda", "CloudFront → WAF → API GW / ALB", "同期", "認証実装チェック Lambda", "endpoint ごとに 1 回",
     "認証情報なしでリクエストし、正しく拒否されるか（401/403、Cookie 系は 302）を確認する。2xx なら認証漏れ。",
     "認証不要（Public 経路）/ HTTPS 443 / X-Auth-Probe ヘッダ付与（WAF 誤検知回避）",
     "チェック Lambda -> CloudFront : GET（認証ヘッダなし, X-Auth-Probe）\nCloudFront -> API GW : Origin Protection 付与\nAPI GW --> チェック Lambda : ステータスコード",
     "公開印（x-synthetics-skip-auth-check）付き endpoint の扱い、タイムアウト値、endpoint 単位の継続（1 件失敗で打ち切らない）を明記する"),
    ("P-27", "正規検査", "正規検査（Positive probe）", "B 検査",
     "認証実装チェック Lambda", "CloudFront → WAF → API GW / ALB", "同期", "認証実装チェック Lambda", "endpoint ごとに 1 回",
     "有効なトークン付きでリクエストし、200 が返ることで API 稼働と検査自体の健全性を確認する。",
     "Bearer（P-25 のトークン）/ HTTPS 443 / X-Auth-Probe ヘッダ付与",
     "チェック Lambda -> CloudFront : GET（Authorization: Bearer …, X-Auth-Probe）\nAPI GW --> チェック Lambda : ステータスコード",
     "参照系のみ実行すること（更新系を叩かない）、対象 endpoint の選定規則を明記する"),
    ("P-28", "判定・分類", "判定・分類（4×4）", "B 検査",
     "認証実装チェック Lambda（内部）", "—", "内部", "認証実装チェック Lambda", "endpoint ごとに 1 回",
     "Negative/Positive の結果の組み合わせから、CRITICAL / WARN / INFO / OK を判定する。",
     "—（AWS 呼び出しなし）",
     "チェック Lambda : (Negative, Positive) -> 4×4 真偽値表 -> severity 決定\nWAF 起因の 403 は WARN（構成）に分類",
     "4×4 の全組み合わせと severity の対応表を貼る（code-samples/README §2 が正）"),
    ("P-29", "検査メトリクス", "検査結果メトリクス出力", "B 検査",
     "認証実装チェック Lambda", "CloudWatch", "同期", "認証実装チェック Lambda", "検査ごとに 1 回",
     "検査結果をメトリクス化し、保険系アラーム（AuthCheckCritical > 0）の入力とする。",
     "cloudwatch:PutMetricData / monitoring.{region}.amazonaws.com",
     "チェック Lambda -> CloudWatch : PutMetricData(AuthCheckCritical ほか, ディメンション=appId/env)",
     "ディメンション設計（アプリ数比例でメトリクス課金が増える点に注意）を明記する"),
    ("P-30", "検知イベント送付", "検知イベント送付", "B 検査",
     "認証実装チェック Lambda", "Alert Router Lambda", "非同期", "認証実装チェック Lambda", "severity≠OK 時",
     "分類済みの検知イベントを通知の振り分け役へ引き渡す。",
     "lambda:InvokeFunction（InvocationType=Event）",
     "チェック Lambda -> Alert Router : Invoke(Event, 4×4 分類済みイベント)",
     "イベント形式（README §2.6）と、配列でのバッチ送付の可否を明記する"),

    ("P-31", "通知先解決", "通知先解決", "C 通知",
     "Alert Router Lambda", "Monitoring Registry S3（registry/）", "同期", "Alert Router Lambda", "イベントごとに 1 回",
     "台帳の alertRouting から通知先 SNS ARN を解決する。未設定時は全社デフォルトを使う。",
     "s3:GetObject",
     "Alert Router -> registry/ : GetObject（alertRouting 参照）\n未設定 -> 全社デフォルト ARN",
     "2 段解決（アプリ個別 → 全社デフォルト）の順序と、未解決時に throw する仕様を明記する"),
    ("P-32", "通知送信", "通知送信", "C 通知",
     "Alert Router Lambda", "SNS（P1 Security / P2 Platform / P3 App）", "同期", "Alert Router Lambda", "イベントごとに 1 回",
     "severity に応じた宛先へ通知を送信する。",
     "sns:Publish / sns.{region}.amazonaws.com",
     "Alert Router -> SNS(P1|P2|P3) : Publish（件名・本文・appId・endpoint・判定根拠）",
     "通知本文のテンプレートと、1 件でも失敗したら throw（DLQ 発火）する仕様を明記する"),
    ("P-33", "保険系アラーム", "保険系アラーム発報", "C 通知",
     "CloudWatch Alarm", "SNS", "非同期", "CloudWatch", "閾値超過時",
     "Alert Router 経路とは独立に、メトリクス閾値からも発報する（発報 2 系統の保険側）。",
     "CloudWatch Alarm → SNS（アラームアクション）",
     "CloudWatch Alarm(AuthCheckCritical > 0) -> SNS : 通知",
     "即時系（Alert Router）と保険系の重複通知の扱いを明記する"),

    ("P-40", "メタ監視発報", "メタ監視発報（MM-1〜5）", "D 運用・異常",
     "CloudWatch Alarm", "SNS（P2 Platform）", "非同期", "CloudWatch", "閾値超過時",
     "監視系そのものの停止・失敗を検知して発報する（監視の空白＝検知の空白を防ぐ）。",
     "CloudWatch Alarm → SNS",
     "Alarm(DiscoveryLastSuccess 2h 欠損 / Lambda Errors / DLQ 滞留) -> SNS(P2)",
     "MM-1〜5 の各アラーム定義（対象メトリクス・閾値・評価期間）を一覧化する"),
    ("P-41", "DLQ 退避・再処理", "DLQ 退避・再処理", "D 運用・異常",
     "Lambda（非同期実行の失敗）", "SQS DLQ → 運用者", "非同期", "Lambda / 運用者", "失敗時",
     "非同期呼び出しがリトライ後も失敗した場合に退避し、後から再処理・原因調査できるようにする。",
     "Lambda 非同期呼び出しの DLQ 設定 / sqs:ReceiveMessage・DeleteMessage（運用者）",
     "Lambda(失敗) -> 自動リトライ(2 回) -> SQS DLQ\n運用者 -> DLQ : 内容確認 -> 再実行",
     "DLQ の保持期間・再処理手順（Runbook 化）・滞留アラーム（MM-4/5）を明記する"),
    ("P-42", "台帳の手動更新", "台帳の手動更新", "D 運用・異常",
     "運用者", "Monitoring Registry S3（registry/）", "同期", "運用者", "随時",
     "enabled の切替や alertRouting の設定など、中央管理項目を運用者が更新する。",
     "s3:GetObject / s3:PutObject（ETag 条件付き）",
     "運用者 -> registry/ : GetObject -> 編集 -> PutObject（If-Match）\n巡回と競合した場合は 412 -> 再取得",
     "巡回（P-09）との競合手順、変更履歴の追跡（Versioning）、承認フローの要否を明記する"),

    ("P-50", "資材アップロード", "監視資材アップロード", "E アプリ接点",
     "ベンダー CI（デプロイパイプライン最終段）", "App アカウントの資材バケット", "同期", "ベンダー CI", "デプロイのたび",
     "デプロイ成功後に監視資材をアップロードする。本監視の入口であり、アプリ（ベンダー）側の責務。",
     "ArtifactUploadRole-{appId} を Assume / s3:PutObject（{appId}/* 限定）/ s3.{region}.amazonaws.com",
     "ベンダー CI -> STS : AssumeRole(ArtifactUploadRole-{appId})\nベンダー CI -> 資材バケット : PutObject(monitoring.yaml, openapi.yaml, deploy-info.json)",
     "デプロイ成功後に実行すること（順序逆転の禁止）、アップロード漏れは原則アプリ責任（M-Q-17-7）である旨を明記する"),
]

SEC_TABLES = {
    "2. 入力（I）": ["項目", "型", "必須", "取得元", "説明・例"],
    "3. 出力（O）": ["項目", "型", "出力先", "説明・例"],
    "4. 処理手順": ["#", "処理内容", "補足・参照"],
    "6. 例外・異常系": ["ケース", "検知方法", "動作", "通知", "参照"],
    "9. ログ・メトリクス": ["種別", "名称・項目", "内容", "備考"],
    "11. 未決事項": ["ID", "内容", "確認先", "期限"],
}


def style_title(ws, text, span=5):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
    c = ws.cell(row=1, column=1, value=text)
    c.font, c.fill, c.alignment = TITLE, TITLE_FILL, Alignment(vertical="center")
    ws.row_dimensions[1].height = 26


def section(ws, row, text, span=5):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row=row, column=1, value=text)
    c.font, c.fill = SEC, SEC_FILL
    return row + 1


def kv(ws, row, label, value, todo=False, span=5):
    lc = ws.cell(row=row, column=1, value=label)
    lc.font, lc.fill, lc.border, lc.alignment = HDR, LBL_FILL, BOX, WRAP
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=span)
    vc = ws.cell(row=row, column=2, value=value)
    vc.font, vc.border, vc.alignment = BASE, BOX, WRAP
    if todo:
        vc.fill = TODO_FILL
    for col in range(3, span + 1):
        ws.cell(row=row, column=col).border = BOX
    return row + 1


def table(ws, row, headers, blank_rows=3):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font, c.fill, c.border, c.alignment = HDR, HDR_FILL, BOX, WRAP
    for r in range(row + 1, row + 1 + blank_rows):
        for i in range(1, len(headers) + 1):
            c = ws.cell(row=r, column=i)
            c.border, c.fill, c.alignment = BOX, TODO_FILL, WRAP
    return row + 1 + blank_rows


def build_process_sheet(wb, p):
    (pid, short, name, cat, src, dst, mode, actor, freq, summary, perm, seq, hint) = p
    ws = wb.create_sheet(f"{pid}_{short}"[:31])
    ws.sheet_properties.tabColor = {"A": "FFD966", "B": "9DC3E6", "C": "F4B183",
                                    "D": "C9C9C9", "E": "A9D18E"}[cat[0]]
    for col, w in zip("ABCDE", (22, 26, 20, 30, 46)):
        ws.column_dimensions[col].width = w

    style_title(ws, f"{pid}  {name}")
    r = 2
    r = section(ws, r, "0. 概要")
    r = kv(ws, r, "概要", summary)
    r = section(ws, r, "1. 基本情報")
    r = kv(ws, r, "区分", cat)
    r = kv(ws, r, "フロー内の位置", "", todo=True)
    r = kv(ws, r, "起動元", src)
    r = kv(ws, r, "起動先", dst)
    r = kv(ws, r, "呼び出し方式", mode)
    r = kv(ws, r, "実行主体", actor)
    r = kv(ws, r, "頻度・タイミング", freq)
    r = kv(ws, r, "前提条件・事前状態", "", todo=True)
    r += 1
    for title, headers in list(SEC_TABLES.items())[:2]:
        r = section(ws, r, title)
        r = table(ws, r, headers)
        r += 1
    r = section(ws, r, "4. 処理手順")
    r = table(ws, r, SEC_TABLES["4. 処理手順"], blank_rows=0)
    c = ws.cell(row=r, column=1, value=1)
    c.border, c.fill, c.alignment = BOX, TODO_FILL, WRAP
    c2 = ws.cell(row=r, column=2, value="")
    c2.border, c2.fill, c2.alignment = BOX, TODO_FILL, WRAP
    c3 = ws.cell(row=r, column=3, value=f"【設計時の要記載】{hint}")
    c3.border, c3.fill, c3.alignment = BOX, TODO_FILL, WRAP
    r += 1
    for extra in range(3):
        for i in range(1, 4):
            cc = ws.cell(row=r + extra, column=i)
            cc.border, cc.fill, cc.alignment = BOX, TODO_FILL, WRAP
    r += 3
    r += 1
    r = section(ws, r, "5. シーケンス（矢印記法。必要に応じて図を貼付）")
    ws.merge_cells(start_row=r, start_column=1, end_row=r + 4, end_column=5)
    sc = ws.cell(row=r, column=1, value=seq)
    sc.font, sc.border, sc.alignment = BASE, BOX, WRAP
    r += 6
    r = section(ws, r, "6. 例外・異常系")
    r = table(ws, r, SEC_TABLES["6. 例外・異常系"])
    r += 1
    r = section(ws, r, "7. 冪等性・リトライ")
    r = kv(ws, r, "再実行時の振る舞い", "", todo=True)
    r = kv(ws, r, "状態更新のタイミング", "", todo=True)
    r += 1
    r = section(ws, r, "8. 権限・エンドポイント")
    r = kv(ws, r, "権限 / エンドポイント", perm)
    r = kv(ws, r, "認証方式・補足", "", todo=True)
    r += 1
    r = section(ws, r, "9. ログ・メトリクス")
    r = table(ws, r, SEC_TABLES["9. ログ・メトリクス"], blank_rows=2)
    r += 1
    r = section(ws, r, "10. 性能・上限")
    r = kv(ws, r, "想定件数 / 所要時間", "", todo=True)
    r = kv(ws, r, "上限・制約", "", todo=True)
    r += 1
    r = section(ws, r, "11. 未決事項")
    table(ws, r, SEC_TABLES["11. 未決事項"], blank_rows=2)
    return ws.title


def build_common_sheets(wb, sheet_titles):
    # 00 表紙
    ws = wb.create_sheet("00_表紙・改訂履歴")
    for col, w in zip("ABCDE", (18, 30, 18, 30, 40)):
        ws.column_dimensions[col].width = w
    style_title(ws, "API 認証実装確認処理　処理設計書")
    r = 3
    for label in ("システム名", "文書名", "版", "作成日", "作成者", "承認者", "参照設計書（md）", "参照コミット"):
        default = {"システム名": "API プラットフォーム / 認証実装確認処理",
                   "文書名": "処理設計書（処理単位の I/O・シーケンス）",
                   "参照設計書（md）": "doc/api-platform/basic-design/10〜18 章"}.get(label, "")
        r = kv(ws, r, label, default, todo=not default)
    r += 1
    r = section(ws, r, "改訂履歴")
    table(ws, r, ["版", "日付", "改訂内容", "作成者", "承認者"], blank_rows=5)

    # 01 位置づけ・前提
    ws = wb.create_sheet("01_位置づけ・前提")
    for col, w in zip("ABCDE", (22, 30, 24, 30, 40)):
        ws.column_dimensions[col].width = w
    style_title(ws, "本書の位置づけ・前提")
    r = 2
    r = section(ws, r, "1. 目的と読者")
    r = kv(ws, r, "目的", "10〜18 章の設計を処理単位に分解し、実装・単体テスト・運用手順の起点とする")
    r = kv(ws, r, "読者", "実装担当 / テスト担当 / 運用担当 / レビュア")
    r = kv(ws, r, "上位設計との関係", "設計の正は md（10〜18 章）。本書と矛盾した場合は md が優先する")
    r += 1
    r = section(ws, r, "2. 前提条件")
    r = table(ws, r, ["ID", "前提", "根拠・参照", "確定状況", "備考"], blank_rows=0)
    presets = [
        ("前提1", "実行基盤は Lambda 3 本（発見 / 認証実装チェック / Alert Router）、すべて VPC 外配置", "10 §10.1.6 D-M-10-4", "確定"),
        ("前提2", "変更検知は App アカウントの S3 監視資材の VersionId 比較（資材オンリー原則）", "17 §17.2 / ADR-061 追記 2026-08-21", "確定"),
        ("前提3", "実行モードは 2 つ（自動差分検査＝モード1 / 全量検査＝モード2：日次定期＋手動）", "18 §18.1", "確定"),
        ("前提4", "台帳・spec は共通基盤アカウントの Monitoring Registry（S3 ×1）に集約", "12 / 13 章", "確定"),
        ("前提5", "資材のアップロード漏れ・内容誤りは原則アプリ（ベンダー）責任", "17 §17.2.2 / D-M-17-8（顧客合意 M-Q-17-7）", "合意待ち"),
        ("前提6", "検査は実利用者と同じ経路（CloudFront + WAF）を通過する", "10 §10.1.6 経路 A", "確定"),
    ]
    for i, (pid, txt, ref, st) in enumerate(presets):
        for j, v in enumerate((pid, txt, ref, st, "")):
            c = ws.cell(row=r + i, column=j + 1, value=v)
            c.font, c.border, c.alignment = BASE, BOX, WRAP
    r += len(presets) + 1
    r = section(ws, r, "3. 用語")
    r = table(ws, r, ["用語", "意味", "参照", "", ""], blank_rows=4)

    # 02 全体構成
    ws = wb.create_sheet("02_全体構成")
    for col, w in zip("ABCDE", (26, 20, 46, 24, 20)):
        ws.column_dimensions[col].width = w
    style_title(ws, "全体構成（アカウント・リソース・通信経路）")
    r = 2
    r = section(ws, r, "1. リソース一覧（10 §10.1.5 より転記）")
    r = table(ws, r, ["リソース", "サービス", "役割", "配置アカウント", "詳細章"], blank_rows=10)
    r += 1
    r = section(ws, r, "2. 通信経路（10 §10.1.6 より転記）")
    r = table(ws, r, ["#", "経路", "中身", "通る境界", "必要な許可"], blank_rows=5)
    r += 1
    r = section(ws, r, "3. 構成図（貼付領域）")
    ws.merge_cells(start_row=r, start_column=1, end_row=r + 14, end_column=5)
    c = ws.cell(row=r, column=1, value="※ 10 章 §10.1.6 の AWS リソース構成図を貼り付ける")
    c.font, c.border, c.alignment = BASE, BOX, WRAP

    # 03 処理一覧
    ws = wb.create_sheet("03_処理一覧")
    for col, w in zip("ABCDEFGH", (10, 26, 16, 30, 30, 12, 40, 14)):
        ws.column_dimensions[col].width = w
    style_title(ws, "処理一覧（本書のハブ）", span=8)
    r = 3
    headers = ["処理ID", "処理名", "区分", "起動元", "起動先", "方式", "概要", "シート"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=r, column=i, value=h)
        c.font, c.fill, c.border, c.alignment = HDR, HDR_FILL, BOX, WRAP
    ws.freeze_panes = ws.cell(row=r + 1, column=1)
    for i, p in enumerate(P):
        rr = r + 1 + i
        vals = (p[0], p[2], p[3], p[4], p[5], p[6], p[9])
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=rr, column=j, value=v)
            c.font, c.border, c.alignment = BASE, BOX, WRAP
        title = sheet_titles[i]
        c = ws.cell(row=rr, column=8, value=title)
        c.hyperlink = f"#'{title}'!A1"
        c.font, c.border, c.alignment = LINK, BOX, WRAP
    ws.auto_filter.ref = f"A{r}:H{r + len(P)}"

    # 04 共通仕様
    ws = wb.create_sheet("04_共通仕様")
    for col, w in zip("ABCDE", (24, 60, 30, 20, 20)):
        ws.column_dimensions[col].width = w
    style_title(ws, "共通仕様（全処理に適用）")
    r = 2
    r = section(ws, r, "1. 全処理共通の規約")
    commons = [
        ("命名規約", "", "04 章 / 組織標準"),
        ("ログ出力", "相関 ID を必ず出力。トークン・資格情報はマスクする", "06 章 OBS-1〜4"),
        ("ログ保持期間", "", "WBS A14-i"),
        ("エラー処理の原則", "アカウント単位・endpoint 単位で try-catch し、1 件の失敗で全体を止めない", "18 §18.5.2"),
        ("リトライ", "非同期呼び出しは Lambda 標準リトライ（2 回）+ DLQ", "18 §18.5.2"),
        ("冪等性", "at-least-once。状態（lastArtifactVersions）は後続処理の成功後にのみ更新", "18 §18.5.2"),
        ("タイムアウト", "", "設計時に確定"),
        ("環境変数", "", "設計時に確定"),
        ("タグ", "app-id / env / cost-center / owner を必須付与", "03 章 BL-1"),
        ("宛先 allowlist", "外向き通信の宛先は台帳の baseUrl と設定済み token URL のみ（コードで強制）", "10 §10.1.6 代償統制"),
    ]
    r = table(ws, r, ["項目", "規約", "参照", "", ""], blank_rows=0)
    for i, (k, v, ref) in enumerate(commons):
        for j, val in enumerate((k, v, ref, "", "")):
            c = ws.cell(row=r + i, column=j + 1, value=val)
            c.font, c.border, c.alignment = BASE, BOX, WRAP
            if j == 1 and not val:
                c.fill = TODO_FILL

    # 05 データ定義
    ws = wb.create_sheet("05_データ定義")
    for col, w in zip("ABCDE", (24, 14, 12, 46, 30)):
        ws.column_dimensions[col].width = w
    style_title(ws, "データ定義")
    r = 2
    for title in ("1. 台帳スキーマ（registry/{appId}/{env}.json）",
                  "2. 監視資材（monitoring.yaml / openapi.yaml / deploy-info.json）",
                  "3. イベント payload（Lambda 間）",
                  "4. メトリクス"):
        r = section(ws, r, title)
        r = table(ws, r, ["項目", "型", "必須", "説明・例", "出典"], blank_rows=5)
        r += 1

    # 06 IAM
    ws = wb.create_sheet("06_IAM・権限一覧")
    for col, w in zip("ABCDE", (30, 24, 46, 24, 24)):
        ws.column_dimensions[col].width = w
    style_title(ws, "IAM・権限一覧")
    r = 2
    r = section(ws, r, "1. ロール一覧")
    r = table(ws, r, ["ロール", "使い手", "権限", "配置アカウント", "利用する処理（P-xx）"], blank_rows=8)
    r += 1
    r = section(ws, r, "2. 信頼関係")
    table(ws, r, ["ロール", "信頼元", "条件（ExternalId 等）", "参照", ""], blank_rows=5)

    # 07 非機能
    ws = wb.create_sheet("07_非機能")
    for col, w in zip("ABCDE", (26, 46, 24, 24, 24)):
        ws.column_dimensions[col].width = w
    style_title(ws, "非機能（性能・上限・コスト・監視）")
    r = 2
    for title, headers in (
        ("1. 性能・上限", ["項目", "値・方針", "根拠", "備考", ""]),
        ("2. 監視・アラーム", ["ID", "検知対象", "手段", "閾値", "通知先"]),
        ("3. コスト", ["項目", "算定", "月額", "負担アカウント", "備考"]),
        ("4. 保持期間", ["対象", "保持期間", "根拠", "備考", ""]),
    ):
        r = section(ws, r, title)
        r = table(ws, r, headers, blank_rows=5)
        r += 1


def main():
    wb = Workbook()
    wb.remove(wb.active)
    titles = [build_process_sheet(wb, p) for p in P]
    build_common_sheets(wb, titles)
    order = ["00_表紙・改訂履歴", "01_位置づけ・前提", "02_全体構成", "03_処理一覧",
             "04_共通仕様", "05_データ定義", "06_IAM・権限一覧", "07_非機能"] + titles
    wb._sheets = [wb[t] for t in order]
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT} ({len(wb.sheetnames)} sheets)")


if __name__ == "__main__":
    main()
